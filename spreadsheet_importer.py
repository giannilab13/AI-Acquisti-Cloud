import csv
import hashlib
import io
from pathlib import Path

from openpyxl import load_workbook

from db import get_connection, next_text_id, normalize_basic


HEADER_ALIASES = {
    "code": {
        "CODICE", "CODICE ARTICOLO", "COD ARTICOLO", "COD", "SKU", "ITEM CODE",
        "ARTICLE CODE", "CODICE PRODOTTO", "COD PRODOTTO"
    },
    "description": {
        "DESCRIZIONE", "DESCRIZIONE ARTICOLO", "PRODOTTO", "ARTICOLO",
        "PRODUCT", "DESCRIPTION", "ITEM", "NOME PRODOTTO"
    },
    "unit": {
        "UM", "U M", "UNITA", "UNITA DI MISURA", "UNIT", "UOM", "MISURA"
    },
    "price": {
        "PREZZO", "PREZZO LISTINO", "LISTINO", "PREZZO UNITARIO",
        "PRICE", "UNIT PRICE", "COSTO", "COSTO UNITARIO", "PREZZO EURO"
    },
    "discount_1": {
        "SCONTO", "SCONTO 1", "SCONTO1", "DISC 1", "DISCOUNT 1",
        "SCONTO A", "SCONTO PRINCIPALE"
    },
    "discount_2": {
        "SCONTO 2", "SCONTO2", "DISC 2", "DISCOUNT 2", "SCONTO B"
    },
}


def _normalize_header(value):
    return normalize_basic(str(value or ""))


def _to_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("€", "").replace(" ", "")
    # Italian number support: 1.234,56 -> 1234.56
    if "," in text and "." in text and text.rfind(",") > text.rfind("."):
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except Exception:
        return None


def _discount_value(value):
    """Converte lo sconto in percentuale, anche dal formato percentuale Excel."""
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        number = float(value)
        if 0 < number < 1:
            return round(number * 100, 4)
        return round(number, 4)

    text = str(value).strip().replace("%", "").replace(" ", "")
    if "," in text:
        text = text.replace(",", ".")
    try:
        number = float(text)
        if 0 < number < 1:
            return round(number * 100, 4)
        return round(number, 4)
    except Exception:
        return 0.0


def _net_price(list_price, discount_1=0.0, discount_2=0.0):
    d1 = max(0.0, min(float(discount_1 or 0), 100.0))
    d2 = max(0.0, min(float(discount_2 or 0), 100.0))
    return round(float(list_price) * (1 - d1 / 100.0) * (1 - d2 / 100.0), 4)


def _detect_columns(headers):
    normalized = [_normalize_header(h) for h in headers]
    mapping = {}

    for target, aliases in HEADER_ALIASES.items():
        alias_norm = {_normalize_header(a) for a in aliases}
        for idx, header in enumerate(normalized):
            if header in alias_norm:
                mapping[target] = idx
                break

    if "description" not in mapping:
        raise ValueError(
            "Non trovo una colonna descrizione. Usa un'intestazione come "
            "'Descrizione', 'Prodotto' o 'Articolo'."
        )
    if "price" not in mapping:
        raise ValueError(
            "Non trovo una colonna prezzo. Usa un'intestazione come "
            "'Prezzo', 'Prezzo unitario' o 'Costo'."
        )
    return mapping


def _read_csv(file_bytes):
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Codifica CSV non riconosciuta.")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\\t,")
        delimiter = dialect.delimiter
    except Exception:
        delimiter = ";"

    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return rows


def _read_xlsx(file_bytes):
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    wb.close()
    return rows


def _find_header_row(rows):
    # Cerca nelle prime 15 righe quella che contiene descrizione + prezzo
    for i, row in enumerate(rows[:15]):
        try:
            _detect_columns(row)
            return i
        except Exception:
            pass
    raise ValueError(
        "Non riesco a individuare la riga delle intestazioni. "
        "Il file deve contenere almeno una colonna Descrizione/Prodotto e una colonna Prezzo."
    )


def _ensure_product(conn, supplier_id, code, description, unit):
    code = (str(code).strip() if code is not None else "")
    description = str(description or "").strip()
    normalized = normalize_basic(description)

    if code:
        row = conn.execute(
            """
            SELECT product_id FROM supplier_products
            WHERE supplier_id=? AND supplier_code=?
            """,
            (supplier_id, code),
        ).fetchone()
        if row:
            return row["product_id"], "CODICE_FORNITORE", False

    row = conn.execute(
        """
        SELECT id FROM products
        WHERE UPPER(normalized_name)=UPPER(?)
           OR UPPER(name)=UPPER(?)
        ORDER BY id LIMIT 1
        """,
        (normalized, description),
    ).fetchone()

    if row:
        product_id = row["id"]
        method = "DESCRIZIONE_ESATTA"
        created = False
    else:
        product_id = next_text_id(conn, "products", "PRD")
        conn.execute(
            """
            INSERT INTO products(id,name,normalized_name,category,unit)
            VALUES(?,?,?,?,?)
            """,
            (product_id, description, normalized, "Da classificare", unit or "pz"),
        )
        method = "NUOVO_PRODOTTO"
        created = True

    if code:
        conn.execute(
            """
            INSERT OR IGNORE INTO supplier_products(
                supplier_id,supplier_code,original_description,product_id
            ) VALUES(?,?,?,?)
            """,
            (supplier_id, code, description, product_id),
        )

    return product_id, method, created


def import_price_list_file(file_name, file_bytes, supplier_id):
    conn = get_connection()

    supplier = conn.execute(
        "SELECT id,name FROM suppliers WHERE id=?",
        (supplier_id,),
    ).fetchone()
    if not supplier:
        conn.close()
        raise ValueError("Fornitore non valido.")

    file_hash = hashlib.sha256(file_bytes).hexdigest()
    duplicate = conn.execute(
        "SELECT id FROM price_lists WHERE file_hash=?",
        (file_hash,),
    ).fetchone()
    if duplicate:
        conn.close()
        return {
            "ok": False,
            "duplicate": True,
            "message": f"Listino già importato (ID {duplicate['id']}).",
        }

    extension = Path(file_name).suffix.lower()
    rows = _read_csv(file_bytes) if extension == ".csv" else _read_xlsx(file_bytes)

    if not rows:
        conn.close()
        raise ValueError("Il file è vuoto.")

    header_index = _find_header_row(rows)
    headers = rows[header_index]
    mapping = _detect_columns(headers)

    parsed = []
    skipped = 0

    for row_num, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        # Pad row to header length
        row = list(row) + [None] * max(0, len(headers) - len(row))
        description = row[mapping["description"]] if mapping["description"] < len(row) else None
        price_raw = row[mapping["price"]] if mapping["price"] < len(row) else None

        description = str(description or "").strip()
        price = _to_float(price_raw)

        if not description and price is None:
            continue
        if not description or price is None:
            skipped += 1
            continue

        code = row[mapping["code"]] if "code" in mapping and mapping["code"] < len(row) else ""
        unit = row[mapping["unit"]] if "unit" in mapping and mapping["unit"] < len(row) else "pz"
        unit = str(unit or "pz").strip()

        discount_1 = _discount_value(
            row[mapping["discount_1"]]
            if "discount_1" in mapping and mapping["discount_1"] < len(row)
            else 0
        )
        discount_2 = _discount_value(
            row[mapping["discount_2"]]
            if "discount_2" in mapping and mapping["discount_2"] < len(row)
            else 0
        )
        net_price = _net_price(price, discount_1, discount_2)

        parsed.append({
            "row_number": row_num,
            "code": str(code or "").strip(),
            "description": description,
            "unit": unit,
            "list_price": price,
            "discount_1": discount_1,
            "discount_2": discount_2,
            "net_price": net_price,
        })

    if not parsed:
        conn.close()
        raise ValueError("Non ho trovato righe prodotto valide nel listino.")

    cur = conn.execute(
        """
        INSERT INTO price_lists(
            supplier_id,file_name,file_hash,status,rows_count
        ) VALUES(?,?,?,?,?)
        """,
        (supplier_id, file_name, file_hash, "IMPORTATO", len(parsed)),
    )
    price_list_id = cur.lastrowid

    created_products = 0
    matched_products = 0
    new_products = 0

    preview = []
    for item in parsed:
        product_id, method, created = _ensure_product(
            conn,
            supplier_id,
            item["code"],
            item["description"],
            item["unit"],
        )
        if created:
            created_products += 1
            new_products += 1
        else:
            matched_products += 1

        conn.execute(
            """
            INSERT INTO price_list_lines(
                price_list_id,row_number,supplier_code,original_description,
                product_id,match_method,unit,offered_price,
                list_price,discount_1,discount_2,net_price
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                price_list_id,item["row_number"],item["code"],item["description"],
                product_id,method,item["unit"],item["net_price"],
                item["list_price"],item["discount_1"],item["discount_2"],item["net_price"],
            ),
        )

        if len(preview) < 10:
            preview.append({
                **item,
                "product_id": product_id,
                "match_method": method,
            })

    conn.commit()
    conn.close()

    return {
        "ok": True,
        "duplicate": False,
        "price_list_id": price_list_id,
        "supplier": supplier["name"],
        "file_name": file_name,
        "rows_count": len(parsed),
        "skipped_rows": skipped,
        "matched_products": matched_products,
        "created_products": created_products,
        "preview": preview,
    }
